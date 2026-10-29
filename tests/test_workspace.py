import base64
import io
import json
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from server import create_app, digest
from documents import export_document, import_document


class WorkspaceTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.app = create_app(dict(DB_PATH=self.temp.name + '/db', ADMIN_EMAIL='admin@example.com', ADMIN_PASSWORD='long-test-password', COOKIE_SECURE=False, OPENAI_API_KEY='test', TESTING=True), start_worker=False)
        self.c = self.app.test_client(); self.c.post('/api/login', json=dict(email='admin@example.com', password='long-test-password'))
        self.csrf = self.c.get('/api/me').json['csrf']; self.headers = {'X-CSRF-Token': self.csrf}
        self.rid = self.call('/rooms', {'name': 'Docs', 'description': ''}).json['id']
        self.base = '/rooms/' + self.rid
    def tearDown(self): self.temp.cleanup()
    def call(self, path, data, method='POST', client=None, headers=None):
        return (client or self.c).open('/api' + path, json=data, method=method, headers=headers or self.headers)
    def doc(self):
        return self.call(self.base + '/documents', {'title': 'Launch plan', 'body': '## Purpose\nA **clear** plan.\n\n| Task | Owner |\n| --- | --- |\n| Review | James |\n\n- Ship it'}).json['id']
    def outsider(self):
        other = self.call('/rooms', {'name': 'Other', 'description': ''}).json['id']
        invite = self.call('/rooms/' + other + '/invites', {'email': 'm@example.com'}).json
        client = self.app.test_client(); client.post('/api/join', json=dict(email='m@example.com', name='M', password='long-member-password', token=invite['token']))
        return client, {'X-CSRF-Token': client.get('/api/me').json['csrf']}
    def test_real_downloads_and_import_roundtrip(self):
        from docx import Document
        from pypdf import PdfReader
        did = self.doc()
        for fmt in ('docx', 'pdf', 'md', 'txt', 'html'):
            response = self.c.get('/api' + self.base + '/documents/' + did + '/download/' + fmt)
            self.assertEqual(response.status_code, 200); self.assertIn('attachment;', response.headers['Content-Disposition'])
            if fmt == 'docx':
                doc = Document(io.BytesIO(response.data)); self.assertEqual(doc.paragraphs[0].text, 'Launch plan'); self.assertEqual(doc.tables[0].cell(1, 1).text, 'James')
                imported = self.c.post('/api' + self.base + '/documents/import', data=response.data, content_type='application/octet-stream', headers={**self.headers, 'X-File-Name': 'Launch.docx'})
                self.assertEqual(imported.status_code, 201)
            if fmt == 'pdf': self.assertIn('Launch plan', PdfReader(io.BytesIO(response.data)).pages[0].extract_text())
    def test_conflict_history_and_branch_preserve_original(self):
        did = self.doc(); path = self.base + '/documents/' + did
        result = self.call(path, dict(title='Edited', body='New body', version=1), 'PATCH'); self.assertEqual(result.json['version'], 2)
        self.assertEqual(self.call(path, dict(title='Stale', body='Clobber', version=1), 'PATCH').status_code, 409)
        saved = self.c.get('/api' + path).json; self.assertEqual(saved['title'], 'Edited'); self.assertEqual(saved['history'][0]['title'], 'Launch plan')
        branch = self.call(path + '/branch', dict(title='Branch', body='Unsaved draft')).json['id']
        self.assertEqual(self.c.get('/api' + self.base + '/documents/' + branch).json['parent_id'], did)
        self.assertEqual(self.c.get('/api' + path).json['body'], 'New body')
    def test_document_csrf_and_isolation_cover_exports(self):
        did = self.doc(); path = self.base + '/documents/' + did
        outsider, headers = self.outsider()
        self.assertEqual(outsider.get('/api' + path + '/download/docx').status_code, 404)
        self.assertEqual(self.call(path, dict(title='Bad', body='Bad', version=1), 'PATCH', outsider, headers).status_code, 404)
        self.assertEqual(self.c.patch('/api' + path, json=dict(title='Bad', body='Bad', version=1)).status_code, 403)
        self.assertEqual(self.app.test_client().get('/api' + path + '/download/docx').status_code, 401)
    def test_export_escapes_html_and_import_rejects_bad_files(self):
        html = export_document('<script>x</script>', '<img src=x onerror=alert(1)>', 'html').decode()
        self.assertNotIn('<script>', html); self.assertNotIn('<img', html)
        with self.assertRaises(ValueError): import_document('run.exe', b'MZ')
        response = self.c.post('/api' + self.base + '/documents/import', data=b'bad zip', content_type='application/octet-stream', headers={**self.headers, 'X-File-Name': 'bad.docx'})
        self.assertEqual(response.status_code, 400)
    def configure(self, tools):
        self.call(self.base + '/workmate', dict(agent_name='REIGN', prompt='Help.', enabled=True, mode='mention', tools=tools), 'PATCH')
        self.call(self.base + '/messages', dict(content='REIGN help', request_id='test'))
    def test_agent_reads_updates_and_branches_with_version_guard(self):
        did = self.doc(); self.configure(['read_document', 'update_document', 'branch_document'])
        agent = self.app.extensions['agent']
        def provider(payload, key):
            if not any(i.get('type') == 'function_call_output' for i in payload['input']):
                return {'output': [{'type': 'function_call', 'name': 'update_document', 'call_id': 'edit', 'arguments': json.dumps(dict(document_id=did, version=1, title='AI revision', body='Updated by REIGN'))}]}
            return {'output': [{'type': 'message', 'content': [{'type': 'output_text', 'text': 'Saved.'}]}]}
        agent.provider = provider; agent.process_one()
        saved = self.c.get('/api' + self.base + '/documents/' + did).json
        self.assertEqual(saved['version'], 2); self.assertEqual(saved['title'], 'AI revision')
    def test_hosted_images_are_persisted_and_room_scoped(self):
        self.configure(['image_generation', 'code_interpreter', 'web_search'])
        agent = self.app.extensions['agent']; seen = []
        def provider(payload, key):
            seen.append(payload)
            return {'output': [{'type': 'image_generation_call', 'id': 'image', 'result': base64.b64encode(b'png-test-content').decode()}, {'type': 'web_search_call', 'action': {'sources': [{'url': 'https://example.com', 'title': 'Source'}]}}, {'type': 'message', 'content': [{'type': 'output_text', 'text': 'Done.'}]}]}
        agent.provider = provider; agent.process_one()
        room = self.c.get('/api' + self.base).json
        self.assertEqual({t['type'] for t in seen[0]['tools']}, {'code_interpreter', 'image_generation', 'web_search'})
        self.assertEqual(room['messages'][-1]['sources'][0]['url'], 'https://example.com')
        file = room['files'][0]; self.assertEqual(self.c.get('/api' + self.base + '/files/' + file['id']).data, b'png-test-content')
        outsider, _ = self.outsider(); self.assertEqual(outsider.get('/api' + self.base + '/files/' + file['id']).status_code, 404)
    def test_google_setup_is_admin_only_encrypted_and_state_single_use(self):
        outsider, headers = self.outsider()
        self.assertEqual(outsider.get('/api/integrations/google').status_code, 403)
        self.assertFalse(self.c.get('/api/integrations/google').json['configured'])
        payload = dict(client_id='test.apps.googleusercontent.com', client_secret='test-secret-value', origin='https://workmates.example')
        self.assertEqual(self.call('/integrations/google/configure', payload).status_code, 200)
        status = self.c.get('/api/integrations/google').json
        self.assertTrue(status['configured']); self.assertNotIn('test-secret-value', json.dumps(status))
        dbfile = Path(self.temp.name + '/db').read_bytes(); self.assertNotIn(b'test-secret-value', dbfile)
        start = self.call('/integrations/google/connect', {}).json
        from urllib.parse import parse_qs, urlparse
        params = parse_qs(urlparse(start['url']).query); self.assertEqual(params['code_challenge_method'], ['S256'])
        drive = self.app.extensions['drive']; drive.token_request = lambda p: dict(refresh_token='refresh', access_token='access', scope='https://www.googleapis.com/auth/drive')
        drive.http = lambda *a, **k: {'user': {'emailAddress': 'connected@example.com'}}
        cid = drive.complete(params['state'][0], 'code')
        self.assertEqual(drive.connection(cid)['label'], 'connected@example.com')
        with self.assertRaises(ValueError): drive.complete(params['state'][0], 'replayed')
        self.assertEqual(len(self.c.get('/api/integrations/google').json['connections']), 1)
    def setup_mount(self):
        drive = self.app.extensions['drive']; user = self.c.get('/api/me').json['id']
        with self.app.extensions['store'].connect() as db:
            db.execute('INSERT INTO drive_connections VALUES(?,?,?,?,?)', ('connection', user, 'Google', drive.seal({'refresh_token': 'secret'}), 1))
            db.execute('INSERT INTO drive_mounts VALUES(?,?,?,?,?,?)', ('mount', self.rid, 'connection', 'rootfolder', 'Research', 1))
        return drive
    def test_drive_mount_scope_prevents_outside_file_and_shared_drive_flags(self):
        drive = self.setup_mount(); calls = []
        files = {'rootfolder': dict(id='rootfolder', name='Research', mimeType='application/vnd.google-apps.folder'), 'inside': dict(id='inside', name='Brief.txt', mimeType='text/plain', parents=['rootfolder']), 'outside': dict(id='outside', name='Private.txt', mimeType='text/plain', parents=[])}
        def api(cid, path, method='GET', payload=None, params=None, raw=False, media=None):
            calls.append((path, params))
            if params and params.get('alt') == 'media': return b'Inside document'
            if path == 'files': return {'files': [files['inside']]}
            return files[path.split('/')[1]]
        drive.api = api
        response = self.call(self.base + '/drive/browse', dict(mount_id='mount'))
        self.assertEqual(response.status_code, 200); self.assertEqual(calls[-1][1]['includeItemsFromAllDrives'], 'true')
        self.assertEqual(self.call(self.base + '/drive/import', dict(mount_id='mount', file_id='outside')).status_code, 400)
        imported = self.call(self.base + '/drive/import', dict(mount_id='mount', file_id='inside'))
        self.assertEqual(imported.status_code, 201)
        self.assertEqual(self.c.get('/api' + self.base + '/documents/' + imported.json['id']).json['body'], 'Inside document')
        outsider, headers = self.outsider(); self.assertEqual(self.call(self.base + '/drive/browse', dict(mount_id='mount'), client=outsider, headers=headers).status_code, 404)
    def test_drive_save_and_subfolder_branch_use_selected_mount(self):
        drive = self.setup_mount(); did = self.doc(); calls=[]
        drive.within=lambda m, f: dict(id=f, mimeType='application/vnd.google-apps.folder')
        def api(cid, path, method='GET', payload=None, params=None, raw=False, media=None):
            calls.append((cid, payload, params, media)); return dict(id='created', name=payload['name'], webViewLink='https://docs.google.com/document/d/created')
        drive.api=api
        self.assertEqual(self.call(self.base + '/drive/save', dict(mount_id='mount', document_id=did)).status_code, 201)
        self.assertTrue(calls[0][3].startswith(b'PK')); self.assertEqual(calls[0][1]['parents'], ['rootfolder'])
        self.assertEqual(self.call(self.base + '/drive/branch', dict(mount_id='mount', name='Branch')).status_code, 201)
        self.assertEqual(calls[1][1]['mimeType'], 'application/vnd.google-apps.folder')
        self.assertEqual(self.call('/integrations/google/connection', {}, 'DELETE').status_code, 200)
        self.assertEqual(self.c.get('/api' + self.base).json['drive_mounts'], [])


if __name__ == '__main__': unittest.main()
