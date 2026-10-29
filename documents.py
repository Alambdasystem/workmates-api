"""Portable Markdown documents and bounded, non-executable imports."""
import html
import io
import re
import zipfile
from pathlib import Path

MIMES = {'docx': 'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
         'pdf': 'application/pdf', 'md': 'text/markdown; charset=utf-8',
         'txt': 'text/plain; charset=utf-8', 'html': 'text/html; charset=utf-8'}


def blocks(body):
    lines = body.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        i += 1
        if not line.strip():
            continue
        if line.startswith('```'):
            code = []
            while i < len(lines) and not lines[i].startswith('```'):
                code.append(lines[i]); i += 1
            i += 1
            yield 'code', '\n'.join(code)
        elif i < len(lines) and '|' in line and re.fullmatch(r'[\s|:\-]+', lines[i]) and '-' in lines[i]:
            rows = [[x.strip() for x in line.strip('|').split('|')]]
            i += 1
            while i < len(lines) and '|' in lines[i] and lines[i].strip():
                rows.append([x.strip() for x in lines[i].strip('|').split('|')]); i += 1
            width = len(rows[0])
            yield 'table', [(r + [''] * width)[:width] for r in rows]
        elif re.match(r'^#{1,3}\s', line):
            yield 'h' + str(len(line.split(' ')[0])), line.split(' ', 1)[1]
        elif re.match(r'^\s*[-*]\s+', line):
            yield 'bullet', re.sub(r'^\s*[-*]\s+', '', line)
        elif re.match(r'^\s*\d+[.)]\s+', line):
            yield 'number', re.sub(r'^\s*\d+[.)]\s+', '', line)
        else:
            yield 'p', line


def plain(text):
    return re.sub(r'\[([^]]+)\]\(([^)]+)\)', r'\1 (\2)', text).replace('**', '').replace('`', '')


def inline_html(text):
    text = html.escape(text)
    text = re.sub(r'\*\*(.+?)\*\*', r'<b>\1</b>', text)
    text = re.sub(r'`([^`]+)`', r'<code>\1</code>', text)
    return text


def export_document(title, body, fmt):
    if fmt not in MIMES:
        raise ValueError('Choose Word, PDF, Markdown, text, or HTML.')
    if fmt in ('md', 'txt'):
        return (('# ' + title + '\n\n' + body) if fmt == 'md' else (title + '\n\n' + plain(body))).encode()
    if fmt == 'html':
        parts = ['<!doctype html><html><head><meta charset="utf-8"><title>' + html.escape(title) + '</title></head><body><h1>' + html.escape(title) + '</h1>']
        for kind, value in blocks(body):
            if kind == 'table':
                parts.append('<table border="1">' + ''.join('<tr>' + ''.join('<td>' + inline_html(c) + '</td>' for c in r) + '</tr>' for r in value) + '</table>')
            else:
                tag = {'bullet': 'p', 'number': 'p', 'code': 'pre'}.get(kind, kind)
                parts.append('<' + tag + '>' + inline_html(value) + '</' + tag + '>')
        return (''.join(parts) + '</body></html>').encode()
    out = io.BytesIO()
    if fmt == 'docx':
        from docx import Document
        from docx.shared import Inches, Pt, RGBColor
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        doc = Document()
        for section in doc.sections:
            section.top_margin = section.bottom_margin = Inches(.8)
            section.left_margin = section.right_margin = Inches(.85)
        doc.styles['Normal'].font.name = 'Calibri'
        doc.styles['Normal'].font.size = Pt(11)
        for name in ('Title', 'Heading 1', 'Heading 2', 'Heading 3'):
            doc.styles[name].font.color.rgb = RGBColor(0, 0, 0)
            for border in doc.styles[name].element.xpath('./w:pPr/w:pBdr'):
                border.getparent().remove(border)
        doc.add_paragraph(title, 'Title')
        for kind, value in blocks(body):
            if kind == 'table':
                table = doc.add_table(rows=0, cols=len(value[0])); table.style = 'Table Grid'; table.autofit = False
                weights = [max(8, min(45, max(len(r[c]) for r in value))) for c in range(len(value[0]))]
                for column, weight in zip(table.columns, weights):
                    column.width = Inches(6.8 * weight / sum(weights))
                for index, row in enumerate(value):
                    cells = table.add_row().cells
                    for cell, text, weight in zip(cells, row, weights):
                        cell.width = Inches(6.8 * weight / sum(weights))
                        cell.text = plain(text)
                        properties = cell._tc.get_or_add_tcPr()
                        borders = OxmlElement('w:tcBorders')
                        for edge in ('top', 'left', 'bottom', 'right'):
                            border = OxmlElement('w:' + edge); border.set(qn('w:val'), 'single'); border.set(qn('w:sz'), '4'); border.set(qn('w:color'), 'D9D9D9'); borders.append(border)
                        properties.append(borders)
                        margins = OxmlElement('w:tcMar')
                        for edge in ('top', 'left', 'bottom', 'right'):
                            margin = OxmlElement('w:' + edge); margin.set(qn('w:w'), '100'); margin.set(qn('w:type'), 'dxa'); margins.append(margin)
                        properties.append(margins)
                        if index == 0:
                            cell.paragraphs[0].runs[0].bold = True
                            shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), 'E9EDF0'); properties.append(shade)
                repeat = OxmlElement('w:tblHeader'); repeat.set(qn('w:val'), 'true')
                table.rows[0]._tr.get_or_add_trPr().append(repeat)
                doc.add_paragraph()
            else:
                style = {'h1': 'Heading 1', 'h2': 'Heading 2', 'h3': 'Heading 3', 'bullet': 'List Bullet', 'number': 'List Number'}.get(kind)
                para = doc.add_paragraph(style=style)
                for piece in re.split(r'(\*\*.*?\*\*|`[^`]+`)', value):
                    run = para.add_run(plain(piece))
                    run.bold = piece.startswith('**')
                    if kind == 'code' or piece.startswith('`'):
                        run.font.name = 'Consolas'; run.font.size = Pt(9)
        doc.core_properties.author = 'Workmates by Alambda'
        doc.save(out)
    else:
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Preformatted
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.lib import colors
        styles = getSampleStyleSheet()
        story = [Paragraph(html.escape(title), styles['Title']), Spacer(1, 16)]
        for kind, value in blocks(body):
            if kind == 'table':
                width = 468 / len(value[0])
                table = Table([[Paragraph(inline_html(c), styles['BodyText']) for c in r] for r in value], colWidths=[width] * len(value[0]), repeatRows=1, hAlign='LEFT')
                table.setStyle(TableStyle([('GRID', (0, 0), (-1, -1), .5, colors.lightgrey), ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#e9edf0')), ('VALIGN', (0, 0), (-1, -1), 'TOP'), ('TOPPADDING', (0, 0), (-1, -1), 7), ('BOTTOMPADDING', (0, 0), (-1, -1), 7)]))
                story.append(table)
            elif kind == 'code':
                story.append(Preformatted(value, styles['Code'], maxLineLength=85))
            else:
                style = styles[{'h1': 'Heading1', 'h2': 'Heading2', 'h3': 'Heading3'}.get(kind, 'BodyText')]
                story.append(Paragraph(('• ' if kind == 'bullet' else '') + inline_html(value), style))
            story.append(Spacer(1, 8))
        SimpleDocTemplate(out, leftMargin=72, rightMargin=72).build(story)
    return out.getvalue()


def import_document(filename, content):
    ext = Path(filename).suffix.lower()
    if not content or len(content) > 8 * 1024 * 1024:
        raise ValueError('Upload a nonempty file up to 8 MB.')
    if ext in ('.txt', '.md', '.csv', '.json'):
        text = content.decode('utf-8-sig')
    elif ext == '.docx':
        from docx import Document
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            if sum(z.file_size for z in archive.infolist()) > 25 * 1024 * 1024:
                raise ValueError('The expanded document is too large.')
        doc = Document(io.BytesIO(content))
        text = '\n\n'.join(p.text for p in doc.paragraphs)
        for table in doc.tables:
            text += '\n\n' + '\n'.join('| ' + ' | '.join(c.text for c in r.cells) + ' |' for r in table.rows)
    elif ext == '.pdf':
        from pypdf import PdfReader
        pdf = PdfReader(io.BytesIO(content))
        if len(pdf.pages) > 100:
            raise ValueError('Import up to 100 pages at a time.')
        text = '\n\n'.join(p.extract_text() or '' for p in pdf.pages)
    else:
        raise ValueError('Upload Word, text PDF, Markdown, text, CSV, or JSON.')
    if not text.strip():
        raise ValueError('No readable text found. Scanned PDFs need OCR first.')
    if len(text) > 40000:
        raise ValueError('This file has too much text. Import a smaller section (40,000 characters).')
    return text
