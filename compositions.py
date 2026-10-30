"""Deterministic image compositions. No generated code or shell interpolation."""
import io
import math
import subprocess
import tempfile
from pathlib import Path
from PIL import Image, ImageOps

MIMES = {'docx': 'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
         'pptx': 'application/vnd.openxmlformats-officedocument.presentationml.presentation', 'mp4': 'video/mp4'}


def image_data(content):
    try:
        with Image.open(io.BytesIO(content)) as image:
            if image.width * image.height > 25000000 or image.width < 1 or image.height < 1:
                raise ValueError('Use images up to 25 megapixels.')
            image.load()
            if image.format in ('PNG', 'JPEG') and image.getexif().get(274, 1) == 1:
                return content, image.size
            out = io.BytesIO(); ImageOps.exif_transpose(image).convert('RGB').save(out, format='PNG')
            return out.getvalue(), image.size
    except Exception as exc:
        raise ValueError('Use a readable PNG, JPEG, WebP or GIF image for this composition.') from exc


def export_assets(assets):
    """Keep ordinary originals byte-for-byte; bound oversized Office packages."""
    if sum(len(a['content']) for a in assets.values()) <= 24*1024*1024:return assets,False
    copies={}
    for aid,asset in assets.items():
        content,_=image_data(asset['content'])
        with Image.open(io.BytesIO(content)) as image:
            image=ImageOps.exif_transpose(image).convert('RGBA');image.thumbnail((1280,1280))
            canvas=Image.new('RGB',image.size,'white');canvas.paste(image,mask=image.getchannel('A'))
            out=io.BytesIO();canvas.save(out,format='JPEG',quality=88,optimize=True)
            if out.tell()>2*1024*1024:
                out=io.BytesIO();canvas.save(out,format='JPEG',quality=70,optimize=True)
        copies[aid]={**asset,'content':out.getvalue()}
    return copies,True


def make_word(title, introduction, entries):
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    doc = Document(); section = doc.sections[0]
    section.top_margin = section.bottom_margin = Inches(.7)
    section.left_margin = section.right_margin = Inches(.8)
    doc.styles['Normal'].font.name = 'Calibri'; doc.styles['Normal'].font.size = Pt(11)
    doc.styles['Title'].font.color.rgb = RGBColor(0, 0, 0)
    for border in doc.styles['Title'].element.xpath('./w:pPr/w:pBdr'):
        border.getparent().remove(border)
    doc.add_paragraph(title, 'Title')
    if introduction: doc.add_paragraph(introduction)
    for index, (asset, caption) in enumerate(entries):
        if index: doc.add_page_break()
        content, (width, height) = image_data(asset['content'])
        size = min(6.7 / width, 6.7 / height)
        doc.add_picture(io.BytesIO(content), width=Inches(width * size), height=Inches(height * size))
        if caption: doc.add_paragraph(caption, 'Caption')
    out = io.BytesIO(); doc.save(out); return out.getvalue()


def make_presentation(title, slides):
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.oxml.xmlchemy import OxmlElement
    from pptx.enum.shapes import MSO_SHAPE
    prs = Presentation(); prs.slide_width = Inches(13.333333); prs.slide_height = Inches(7.5)
    for spec, asset in slides:
        slide = prs.slides.add_slide(prs.slide_layouts[6]); background = spec['layout'] == 'background' and asset is not None
        if asset:
            content, (width, height) = image_data(asset['content'])
            if background:
                picture = slide.shapes.add_picture(io.BytesIO(content), 0, 0, width=prs.slide_width, height=prs.slide_height)
                ratio = width / height; frame = 16 / 9
                if ratio > frame: picture.crop_left = picture.crop_right = (1 - frame / ratio) / 2
                else: picture.crop_top = picture.crop_bottom = (1 - ratio / frame) / 2
                overlay = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(4.35), prs.slide_width, Inches(3.15))
                overlay.fill.solid(); overlay.fill.fore_color.rgb = RGBColor(10, 30, 24); overlay.line.fill.background()
                alpha = OxmlElement('a:alpha'); alpha.set('val', '72000'); overlay._element.find('.//{http://schemas.openxmlformats.org/drawingml/2006/main}srgbClr').append(alpha)
            else:
                scale = min(7.3 / width, 5.6 / height)
                slide.shapes.add_picture(io.BytesIO(content), Inches(.5), Inches(1.2), width=Inches(width*scale), height=Inches(height*scale))
        x, y, text_width = ((.65, 4.65, 12.0) if background else ((8.25, 1.0, 4.45) if asset else (.8, 1.0, 11.7)))
        head = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(text_width), Inches(1.0)).text_frame
        head.word_wrap = True; head.text = spec['title']
        head.paragraphs[0].font.size = Pt(32); head.paragraphs[0].font.bold = True
        head.paragraphs[0].font.color.rgb = RGBColor(255,255,255) if background else RGBColor(25,65,48)
        body = slide.shapes.add_textbox(Inches(x), Inches(y+1.05), Inches(text_width), Inches(1.55 if background else 4.4)).text_frame
        body.word_wrap = True; body.text = spec['body']
        for para in body.paragraphs:
            para.font.size = Pt(16 if background else 20)
            para.font.color.rgb = RGBColor(255,255,255) if background else RGBColor(55,75,65)
    prs.core_properties.title = title; prs.core_properties.author = 'Workmates by Alambda'
    out = io.BytesIO(); prs.save(out); return out.getvalue()


def make_slideshow(assets, seconds, transition, transition_seconds, check=lambda:None, size=(1280,720), fps=24):
    import imageio_ffmpeg
    if transition not in ('fade','slide','wipe','dissolve'):
        raise ValueError('Choose fade, slide, wipe or dissolve.')
    if not isinstance(seconds, (float,int)) or not math.isfinite(seconds) or not 1 <= seconds <= 8:
        raise ValueError('Use 1–8 seconds per image.')
    if not isinstance(transition_seconds,(float,int)) or not math.isfinite(transition_seconds) or not .2 <= transition_seconds <= min(1,seconds/2):
        raise ValueError('Transition duration must be 0.2–1 second and no more than half the image duration.')
    with tempfile.TemporaryDirectory(prefix='workmates-video-') as folder:
        root = Path(folder); args = [imageio_ffmpeg.get_ffmpeg_exe(), '-hide_banner', '-loglevel', 'error', '-y']
        for index, asset in enumerate(assets):
            check(); content, _ = image_data(asset['content'])
            with Image.open(io.BytesIO(content)) as image:
                frame = ImageOps.fit(image.convert('RGB'), size, method=Image.Resampling.LANCZOS)
                path = root / (str(index)+'.jpg'); frame.save(path, quality=92)
            args += ['-loop','1','-framerate',str(fps),'-t',str(seconds),'-i',str(path)]
        filters = [f'[{i}:v]settb=AVTB,format=yuv420p[v{i}]' for i in range(len(assets))]
        current = 'v0'; effect = {'fade':'fade','slide':'slideleft','wipe':'wipeleft','dissolve':'dissolve'}[transition]
        for index in range(1,len(assets)):
            output='mix'+str(index); offset=index*(seconds-transition_seconds)
            filters.append(f'[{current}][v{index}]xfade=transition={effect}:duration={transition_seconds}:offset={offset}[{output}]');current=output
        duration = len(assets)*seconds - max(0,len(assets)-1)*transition_seconds
        output = root/'slideshow.mp4'
        args += ['-filter_complex_threads','1','-filter_complex',';'.join(filters),'-map','['+current+']','-an','-t',str(duration),'-c:v','libx264','-preset','veryfast','-threads','1','-crf','23','-pix_fmt','yuv420p','-movflags','+faststart',str(output)]
        with tempfile.TemporaryFile() as errors:
            process=subprocess.Popen(args,stdin=subprocess.DEVNULL,stdout=subprocess.DEVNULL,stderr=errors)
            import time
            deadline=time.monotonic()+180
            try:
                while process.poll() is None:
                    check()
                    if time.monotonic()>deadline:raise RuntimeError('Slideshow rendering timed out. Try fewer images or shorter timing.')
                    time.sleep(.2)
                if process.returncode:raise RuntimeError('The slideshow could not be encoded. Try supported images or shorter timing.')
                check();content=output.read_bytes()
                if len(content)>32*1024*1024:raise ValueError('The video is larger than 32 MB. Try shorter timing.')
                return content
            finally:
                if process.poll() is None:process.kill();process.wait()
